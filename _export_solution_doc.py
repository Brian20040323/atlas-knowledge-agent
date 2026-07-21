# -*- coding: utf-8 -*-
"""Generate RAG pain-point solution + prompt spec document."""

from pathlib import Path
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

OUT_DIR = Path(r"c:\Users\czy2004\Desktop")
DOCX = OUT_DIR / "RAG痛点解决方案与Prompt规范.docx"
PDF = OUT_DIR / "RAG痛点解决方案与Prompt规范.pdf"
DOCX_EN = OUT_DIR / "RAG_painpoints_solution_prompts.docx"
PDF_EN = OUT_DIR / "RAG_painpoints_solution_prompts.pdf"
PS1 = Path(r"c:\Users\czy2004\Desktop\AI 全栈开发\_export_solution_pdf.ps1")


def font(run, name="宋体", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Consolas" if name == "Consolas" else "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Consolas" if name == "Consolas" else "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "宋体" if name == "Consolas" else name)
    if color:
        run.font.color.rgb = color


def p(doc, text, *, size=11, bold=False, name="宋体", align=None, after=6, first=False, color=None):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first:
        pf.first_line_indent = Pt(22)
    run = para.add_run(text)
    font(run, name=name, size=size, bold=bold, color=color)
    return para


def h(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 11}
    return p(doc, text, size=sizes.get(level, 11), bold=True, name="黑体", after=8 if level < 3 else 6)


def bullets(doc, items, size=11):
    for t in items:
        para = doc.add_paragraph(style="List Number" if False else "List Bullet")
        para.clear()
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(t)
        font(run, size=size)


def numbered(doc, items, size=11):
    for t in items:
        para = doc.add_paragraph(style="List Number")
        para.clear()
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(t)
        font(run, size=size)


def code_block(doc, text):
    """Monospace-ish paragraph block for prompts."""
    for line in text.strip("\n").split("\n"):
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.15
        pf.left_indent = Cm(0.3)
        run = para.add_run(line if line else " ")
        font(run, name="Consolas", size=9, color=RGBColor(0x1F, 0x29, 0x37))
    doc.add_paragraph()


def add_table(doc, headers, rows, caption=""):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, t in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(t)
        font(run, name="黑体", size=9, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, v in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(v))
            font(run, size=9)
    doc.add_paragraph()
    if caption:
        p(doc, caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    # Cover
    p(doc, "企业制度 / 知识库 RAG Agent", size=12, bold=True, name="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    p(doc, "痛点解决方案与 Prompt 规范", size=20, bold=True, name="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    p(
        doc,
        "对应面试追问 Q1–Q9：多轮记忆 · 检索策略 · 流水线完备性 · 差检索 · 短问澄清 · 拆问并行 · 同义召回 · 空库拒答 · 超时降级",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
    )
    p(doc, "文档版本 v1.0 · 可直接用于方案答辩与工程落地", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

    # 1 Overview
    h(doc, "1. 文档目的与范围", 1)
    p(
        doc,
        "本文针对「制度/知识助手」类 RAG Agent 在真实对话中暴露的九类痛点，给出可落地的解决方案、"
        "系统流程约定，以及可复制到 DeepSeek / GPT 等模型的具体 Prompt。"
        "目标不是堆概念，而是：输入什么场景 → 系统走哪条链路 → 用哪条 Prompt → 输出什么形态。",
        first=True,
    )

    h(doc, "1.1 设计原则（四条硬约束）", 2)
    numbered(
        doc,
        [
            "有据才答：制度数字与条款必须来自检索命中；无命中则拒答，禁止编造。",
            "追问先改写再检索：短句追问不得直接拿原文做向量/切块检索。",
            "缺信息先澄清：极短模糊问优先反问补槽，而不是硬搜硬答。",
            "失败可感知：检索差、空库、超时都必须有明确降级话术与可观测 Trace。",
        ],
    )

    # 2 Pain map
    h(doc, "2. 痛点总览与对策映射", 1)
    add_table(
        doc,
        ["编号", "痛点", "对策模块", "关键 Prompt"],
        [
            ["Q1", "多轮「那薪资呢」丢上下文", "槽位记忆 + 追问改写", "P1 追问改写"],
            ["Q2", "查询策略不清晰", "检索流水线规范", "P1 + P2 + P6"],
            ["Q3", "仅切块相似度不够", "改写→混合召回→过滤→生成", "全链路"],
            ["Q4", "单次检索效果差", "阈值门禁 / 二次改写 / 拒答", "P3 差检索补救"],
            ["Q5", "只问「薪资」太模糊", "澄清反问", "P4 澄清补槽"],
            ["Q6", "复合问题要拆分并行", "问题分解 + 并行检索 + 汇总", "P5 拆问并行"],
            ["Q7", "aaa/bbb 同义漏召回", "同义扩展 + 可选向量", "P2 同义改写"],
            ["Q8", "知识库没有内容", "空库拒答（制度严 / 非制度引导）", "P6 空库拒答"],
            ["Q9", "API 卡住/超时", "超时、降级、熔断", "P7 超时降级"],
        ],
        "表 1　痛点 → 模块 → Prompt 映射",
    )

    # 3 Pipeline
    h(doc, "3. 总体解决架构（推荐流水线）", 1)
    p(doc, "每一轮用户输入，按固定阶段推进（禁止「一句话进、自由发挥出」）：", first=True)
    code_block(
        doc,
        """用户输入
  │
  ├─① 意图路由：闲聊 / 澄清 / 制度LOCAL / 显式联网 / 工具
  ├─② 追问检测：若是追问 → 槽位补全（制度名+主题）
  ├─③ 模糊检测：极短且无槽 → 澄清反问（不检索）
  ├─④ 复合检测：多意图 → 拆子问题（可并行）
  ├─⑤ Query 改写：口语同义 + 制度锚定 + 可选 LLM 改写
  ├─⑥ 检索：切块 + 混合打分（词面为主，向量可选）+ Top-K
  ├─⑦ 质量门禁：弱相关丢弃；不够 → 二次改写或拒答
  ├─⑧ 生成：仅基于 hits；输出答案 + **参考**
  └─⑨ 异常：LLM/外网超时 → 降级话术 / 熔断""",
    )

    h(doc, "3.1 阶段门禁（必须遵守）", 2)
    add_table(
        doc,
        ["阶段", "通过条件", "未通过时"],
        [
            ["意图", "能归入已知 Intent", "默认 LOCAL，禁止自由生成制度条款"],
            ["澄清", "信息足够可检索", "输出澄清问句，本轮不检索"],
            ["检索", "Top-K 且相关分够", "二次改写一次；仍不够则拒答"],
            ["生成", "仅基于 hits（+允许工具）", "不得用模型私货补制度数字"],
            ["引用", "有 document_id / 文档名", "拒答或降级重生成"],
            ["超时", "在超时预算内返回", "P7 降级；外网连续失败熔断"],
        ],
        "表 2　阶段门禁",
    )

    # 4 Solutions per Q
    h(doc, "4. 分痛点解决方案（工程约定）", 1)

    h(doc, "4.1 Q1 多轮记忆：追问必须改写成独立检索句", 2)
    p(doc, "痛点：第二句「那薪资呢？」若直接检索，丢失「员工制度」主题。", first=True)
    p(doc, "方案：", bold=True)
    bullets(
        doc,
        [
            "维护对话槽位 DialogueSlots：policy_name（制度名）、topic（主题）。",
            "识别追问：含「那/呢/还有」且短句，或指代词。",
            "检索前改写：slots + 当前句 →「员工制度 薪资」再检索。",
            "示例：差旅报销 +「那机票呢？」→「差旅报销 机票」。",
        ],
    )
    p(doc, "验收：多轮评测用例「住宿上限 → 那机票呢」必须命中同一制度文档。", first=True)

    h(doc, "4.2 Q2–Q3 查询策略：不止切块算相似度", 2)
    p(doc, "推荐检索策略（默认）：", bold=True)
    numbered(
        doc,
        [
            "规则改写（口语→制度词）+ 可选 LLM 改写（P1/P2）。",
            "文档扫描 → 切块（建议 200–400 字，重叠 10–20%）。",
            "混合召回：词面/TF-IDF 为主；向量可选增强语义。",
            "轻量过滤：低分 / 离题 hit 丢弃（等价轻 rerank）。",
            "生成带引用；无独立重型 rerank 也可先跑通，后期再加 Cross-Encoder。",
        ],
    )
    p(
        doc,
        "结论（答辩口径）：切块+相似度+召回+rerank 是基线，生产还必须有改写、门禁、拒答与可观测。",
        first=True,
    )

    h(doc, "4.3 Q4 检索效果不好：二次改写 → 仍差则拒答", 2)
    bullets(
        doc,
        [
            "第一次 hits 为空或 top_score 低于阈值：触发 P3 生成 2–3 条改写再检索。",
            "制度类仍空：走 P6 拒答，禁止自动外网当制度。",
            "非制度类：可提示上传 /「记住：」/ 显式「联网：」。",
            "Trace 记录 original / rewritten / top_titles / scores，便于复盘。",
        ],
    )

    h(doc, "4.4 Q5 模糊短问：先澄清，不硬搜", 2)
    bullets(
        doc,
        [
            "触发条件：字数≤4 或仅主题词（如「薪资」「报销」）且无历史槽位。",
            "动作：输出 P4 澄清问句（2–4 个选项），本轮不入库检索答案。",
            "用户补全后，进入正常 LOCAL 检索。",
        ],
    )

    h(doc, "4.5 Q6 复合问题：拆子问题 → 并行检索 → 分段汇总", 2)
    bullets(
        doc,
        [
            "用 P5 将长句拆成 2–4 个独立子问题（带完整制度/主题词）。",
            "各子问题并行 search(top_k)，合并去重 hits。",
            "生成时按子问题分小节回答；某子问题无命中则该节拒答，不影响其他节。",
        ],
    )

    h(doc, "4.6 Q7 同义召回：规则表 + LLM 扩展 + 可选向量", 2)
    bullets(
        doc,
        [
            "维护同义词表（酒店→住宿费，校招→校园招聘等）。",
            "检索前用 P2 产出「规范词 + 同义词」查询串。",
            "开启向量后可补语义近邻；规则与向量互补，规则保证关键制度词不丢。",
        ],
    )

    h(doc, "4.7 Q8 空库回答：诚实拒答 + 行动指引", 2)
    bullets(
        doc,
        [
            "制度/政策：P6-A，强调「不能凭推测」，引导补文档或问 HR。",
            "非制度：P6-B，引导导入 / 记住 / 联网（显式）。",
            "禁止话术：一般为、通常是、估计、大概（制度场景）。",
        ],
    )

    h(doc, "4.8 Q9 超时与卡住：预算 + 降级 + 熔断", 2)
    bullets(
        doc,
        [
            "LLM / 外网分别设置超时（如 LLM 60–90s，Web 8–12s）。",
            "LLM 超时：回退已检索 hits 的本地合成，或返回 P7 话术。",
            "外网连续失败 N 次：本轮熔断，跳过后续 web。",
            "对用户可见：说明「超时/降级」，不无限转圈。",
        ],
    )

    # 5 Prompts - the core deliverable
    h(doc, "5. Prompt 规范（可直接落地）", 1)
    p(
        doc,
        "约定：下列 Prompt 中 {{变量}} 由程序填充；输出尽量要求 JSON，便于解析。"
        "模型推荐：DeepSeek V4 Flash（默认）/ Pro（复杂拆问与改写）。",
        first=True,
    )

    # P1
    h(doc, "5.1 P1｜多轮追问改写（解决 Q1/Q2）", 2)
    p(doc, "何时调用：判定为追问，或当前句过短且存在历史槽位。", first=True)
    p(doc, "System", bold=True)
    code_block(
        doc,
        """你是企业知识库检索的「查询改写器」。
任务：把用户的短追问改写成一条独立、可检索的完整问句。
硬性规则：
1) 必须保留上一轮的制度/主题上下文，禁止只输出代词句。
2) 只输出一个改写后的检索查询，不要解释。
3) 不要编造库中不存在的制度名；未知则保留用户原词。
4) 中文输出，长度建议 8–40 字。""",
    )
    p(doc, "User", bold=True)
    code_block(
        doc,
        """对话槽位：
- 制度名：{{policy_name}}
- 主题：{{topic}}

最近用户问：
{{prev_user_question}}

当前用户追问：
{{current_user_question}}

请输出改写后的检索查询（仅一行文本）。""",
    )
    p(doc, "示例", bold=True)
    code_block(
        doc,
        """输入：policy=员工制度，current=那薪资呢？
输出：员工制度 薪资标准与构成

输入：policy=差旅报销，current=那机票呢？
输出：差旅报销 机票报销标准""",
    )

    # P2
    h(doc, "5.2 P2｜同义与口语改写（解决 Q7/Q2）", 2)
    p(doc, "何时调用：检索前；或规则改写后仍命中弱。", first=True)
    p(doc, "System", bold=True)
    code_block(
        doc,
        """你是检索查询扩展器。给定用户问题，输出 JSON：
{
  "canonical_query": "规范检索句（含制度/主题关键词）",
  "synonyms": ["同义词1", "同义词2"],
  "search_queries": ["用于检索的query1", "query2"]
}
规则：
1) search_queries 2–4 条，覆盖口语说法与规范说法。
2) 不要改变用户真实意图；不要引入无关主题。
3) 制度场景优先使用：报销、差旅、考勤、薪资、招聘等规范词。""",
    )
    p(doc, "User", bold=True)
    code_block(
        doc,
        """用户问题：{{user_question}}
已知同义词表（可选）：{{synonym_table_json}}
请只输出 JSON。""",
    )
    p(doc, "示例", bold=True)
    code_block(
        doc,
        """输入：酒店一晚能报多少
输出：
{
  "canonical_query": "差旅报销 住宿费上限",
  "synonyms": ["酒店", "房费", "住宿标准"],
  "search_queries": ["差旅报销 住宿费上限", "差旅 住宿费 一晚", "酒店 报销 额度"]
}""",
    )

    # P3
    h(doc, "5.3 P3｜差检索二次改写（解决 Q4）", 2)
    p(doc, "何时调用：首次检索 hits 为空或 top_score < 阈值。", first=True)
    p(doc, "System", bold=True)
    code_block(
        doc,
        """首次检索效果不佳。请基于失败信息生成最多 3 条「更可检索」的改写查询。
只输出 JSON：{"queries":["...","..."],"reason":"一句话原因"}
要求：
1) 比原句更具体：补制度类型、对象、指标名。
2) 去掉口语虚词、表情、无意义后缀。
3) 若问题本身信息过少，queries 可为空数组，reason 说明应先澄清。""",
    )
    p(doc, "User", bold=True)
    code_block(
        doc,
        """原问题：{{user_question}}
已试查询：{{tried_queries}}
命中标题：{{hit_titles_or_empty}}
最高分：{{top_score}}
请输出 JSON。""",
    )

    # P4
    h(doc, "5.4 P4｜模糊短问澄清（解决 Q5）", 2)
    p(doc, "何时调用：极短/模糊且无足够槽位。本轮不检索作答。", first=True)
    p(doc, "System", bold=True)
    code_block(
        doc,
        """你是企业知识助手的澄清模块。用户问题信息不足，不能直接检索下结论。
输出 JSON：
{
  "need_clarify": true,
  "reply": "给用户看的澄清话术（礼貌、简洁）",
  "options": ["选项1", "选项2", "选项3"]
}
规则：
1) options 2–4 个，覆盖最可能意图。
2) 不要猜测具体制度数字。
3) reply 末尾引导用户点选或补充一句完整问题。""",
    )
    p(doc, "User", bold=True)
    code_block(
        doc,
        """用户只说了：{{user_question}}
可选业务域：制度政策 / 招聘 / 技术知识库 / 其他
请输出 JSON。""",
    )
    p(doc, "示例（用户输入「薪资」）", bold=True)
    code_block(
        doc,
        """{
  "need_clarify": true,
  "reply": "「薪资」范围比较大。你更想了解哪一块？也可以直接说完整一点，例如「正式员工薪资结构」.",
  "options": ["薪资结构/构成", "发薪日与发放规则", "加班费/补贴相关", "校招/社招薪资差异"]
}""",
    )

    # P5
    h(doc, "5.5 P5｜复合问题拆解并行（解决 Q6）", 2)
    p(doc, "何时调用：一句话含 2 个以上独立意图（和/以及/分别/区别）。", first=True)
    p(doc, "System", bold=True)
    code_block(
        doc,
        """你是问题分解器。把用户复杂问拆成可并行检索的子问题。
输出 JSON：
{
  "sub_questions": [
    {"id":"sq1","query":"完整可检索子问题","focus":"主题标签"}
  ],
  "merge_instruction": "汇总时的结构要求"
}
规则：
1) 子问题 2–4 个；每个都必须是自立的完整问句（自带制度/主题词）。
2) 不要拆得过碎；「A和B的区别」可拆成 A定义、B定义、对比点。
3) merge_instruction 要求分段作答，无依据的小节写「材料不足」。""",
    )
    p(doc, "User", bold=True)
    code_block(
        doc,
        """用户问题：{{user_question}}
请输出 JSON。""",
    )
    p(doc, "示例", bold=True)
    code_block(
        doc,
        """输入：薪资和员工制度如何，社招和校招有什么区别
输出：
{
  "sub_questions": [
    {"id":"sq1","query":"员工制度 主要内容与适用范围","focus":"员工制度"},
    {"id":"sq2","query":"薪资制度 结构与发放规则","focus":"薪资"},
    {"id":"sq3","query":"社会招聘与校园招聘 政策区别","focus":"社招校招对比"}
  ],
  "merge_instruction": "分三节：员工制度 / 薪资 / 社招vs校招；无命中的节明确材料不足"
}""",
    )

    h(doc, "5.5.1 汇总 Prompt（拆问后生成）", 3)
    code_block(
        doc,
        """你是企业知识助手。只能依据给定「子问题检索结果」作答，禁止编造制度条款。
按 merge_instruction 分段回答。每节若 hits 为空，写「本节材料不足，不能凭推测」。
文末给出 **参考** 列表（文档名）。不要使用【直接回答】等套话标题。

用户原问题：{{user_question}}
merge_instruction：{{merge_instruction}}

子问题与命中：
{{sub_results_json}}""",
    )

    # P6
    h(doc, "5.6 P6｜空库拒答（解决 Q8）", 2)
    p(doc, "P6-A 制度类（System+固定模板，可不调 LLM）", bold=True)
    code_block(
        doc,
        """目前掌握的材料还不足以完整回答「{{question}}」。

不能凭推测给出制度条款或具体额度。请补充相关制度文档（上传或「记住：标题|内容」），或向 HR/行政确认。""",
    )
    p(doc, "P6-B 非制度类", bold=True)
    code_block(
        doc,
        """目前掌握的材料还不足以完整回答「{{question}}」。

你可以导入相关文件，或用「记住：标题|内容」「自主学习：主题」「联网：主题」补充后再问。""",
    )
    p(doc, "若用 LLM 润色拒答（可选）", bold=True)
    code_block(
        doc,
        """将下面拒答模板改写得更自然，但必须保留：
1) 明确「不足以回答」
2) 明确「不能凭推测」（制度场景）
3) 给出可执行下一步
禁止添加任何具体额度、比例、法条。
模板：{{refuse_template}}
问题：{{question}}""",
    )

    # P7
    h(doc, "5.7 P7｜超时 / 失败降级（解决 Q9）", 2)
    code_block(
        doc,
        """【对用户可见 · 超时降级】
这次生成超时或上游服务暂时不可用。

已为你保留本轮已检索到的依据（如有）：{{hit_titles_or_无}}
建议：稍后重试；若是制度问题，也可先查看已引用文档原文。

【对用户可见 · 外网熔断】
外网检索连续失败，本轮已停止继续联网，避免长时间卡住。
制度类问题请优先依赖本地知识库；如需外网请稍后用「联网：主题」重试。""",
    )

    # P8 answer with citation
    h(doc, "5.8 P8｜有据生成 + 引用（总生成 Prompt）", 2)
    p(doc, "何时调用：LOCAL 且 hits 充足。", first=True)
    code_block(
        doc,
        """你是企业制度/知识助手 Atlas。只能依据「检索片段」回答，禁止使用未出现在片段中的制度数字。
要求：
1) 先用 1–3 句直接回答用户问题，再列必要要点。
2) 使用清晰 Markdown（小标题、列表、加粗）。
3) 文末必须有：
**参考**
- 《文档名》（document_id=...）
4) 禁止【直接回答】【依据】【下一步】等套话标题。
5) 若片段不足以覆盖问题某部分，明确写「材料未覆盖：…」。

用户问题：{{user_question}}
检索片段：
{{hits_context}}""",
    )

    # 6 IO examples
    h(doc, "6. 端到端输入 → 处理 → 输出（对照）", 1)
    add_table(
        doc,
        ["用户输入", "走哪条", "得到什么"],
        [
            ["查询员工制度", "LOCAL 检索", "制度要点 + 参考"],
            ["那薪资呢？", "P1 改写→检索", "员工制度下薪资相关条款 + 参考"],
            ["薪资", "P4 澄清", "反问选项，不编数字"],
            ["薪资和员工制度…社招校招区别", "P5 拆问并行", "分节回答；缺节拒答"],
            ["酒店一晚能报多少", "P2 改写→差旅住宿", "500 元/晚类有据答案（若库中有）"],
            ["打车费能报吗（库无）", "P6 拒答", "不足以…不能凭推测"],
            ["LLM 超时", "P7 降级", "说明超时 + 已有依据/重试建议"],
        ],
        "表 3　端到端对照",
    )

    # 7 Implementation checklist
    h(doc, "7. 工程落地清单", 1)
    numbered(
        doc,
        [
            "实现 DialogueSlots + is_followup + P1 改写接入检索入口。",
            "实现短问澄清门禁（P4），避免「薪资」直接空搜或乱搜。",
            "实现复合问拆解（P5）与并行 search，汇总用分段生成。",
            "同义词表 + P2；向量检索作为可选增强。",
            "统一空库拒答模板 P6；制度场景禁止自动 web。",
            "为 LLM/Web 设超时；失败走 P7；Trace 记录改写与分数。",
            "评测题库覆盖：多轮、短问澄清、拆问、同义、空库、超时模拟。",
        ],
    )

    h(doc, "8. 答辩一句话总结", 1)
    p(
        doc,
        "我们不是「切块算相似度就完事」，而是用槽位记忆与 Prompt 改写解决追问丢上下文，"
        "用澄清/拆问/同义扩展覆盖真实问法，用门禁与拒答保证无据不编造，"
        "用超时降级保证系统卡住时仍可感知、可恢复。",
        first=True,
    )

    p(
        doc,
        "— 完 —  ·  Atlas RAG 痛点解决方案与 Prompt 规范 v1.0",
        size=9,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=0,
    )

    doc.save(str(DOCX))
    # ASCII copy for Word COM
    import shutil

    shutil.copy2(DOCX, DOCX_EN)


def export_pdf():
    PS1.write_text(
        f"""$docx = '{DOCX_EN}'
$pdf = '{PDF_EN}'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$d = $word.Documents.Open($docx, $false, $true)
$d.ExportAsFixedFormat($pdf, 17)
$d.Close($false)
$word.Quit()
""",
        encoding="utf-8",
    )
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(PS1)],
        check=True,
    )
    import shutil

    shutil.copy2(PDF_EN, PDF)


if __name__ == "__main__":
    build()
    export_pdf()
    print("OK", DOCX.exists(), DOCX.stat().st_size, PDF.exists(), PDF.stat().st_size)
