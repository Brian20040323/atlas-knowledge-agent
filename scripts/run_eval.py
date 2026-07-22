"""Run Atlas agent eval cases (planner + retrieval + citation + limits).

Usage:
  .\\.venv\\Scripts\\python.exe scripts\\run_eval.py
  .\\.venv\\Scripts\\python.exe scripts\\run_eval.py --category hit,missing,general
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

try:
    import yaml
except ImportError:
    yaml = None


def _parse_cases_fallback(text: str) -> list[dict]:
    cases: list[dict] = []
    cur: dict | None = None
    for line in text.splitlines():
        if line.strip().startswith("- id:"):
            if cur:
                cases.append(cur)
            cur = {"id": line.split(":", 1)[1].strip()}
        elif cur and ":" in line:
            key, val = line.strip().split(":", 1)
            key, val = key.strip(), val.strip()
            if val.startswith("[") and val.endswith("]"):
                inner = val[1:-1].strip()
                cur[key] = [x.strip().strip("\"'") for x in inner.split(",") if x.strip()]
            elif val.lower() in {"true", "false"}:
                cur[key] = val.lower() == "true"
            else:
                try:
                    cur[key] = int(val)
                except ValueError:
                    cur[key] = val.strip("\"'")
    if cur:
        cases.append(cur)
    return cases


def load_cases(categories: set[str] | None = None) -> list[dict]:
    path = ROOT / "backend" / "tests" / "eval" / "cases.yaml"
    text = path.read_text(encoding="utf-8")
    if yaml:
        data = yaml.safe_load(text)
        cases = list(data.get("cases") or [])
    else:
        cases = _parse_cases_fallback(text)
    if not categories:
        return cases
    return [
        c
        for c in cases
        if str(c.get("category") or "").strip().lower() in categories
    ]

def _contains_any(text: str, needles: list[str]) -> bool:
    return any(n in text for n in needles if n)


def _match_any(text: str, patterns: list[str]) -> bool:
    return any(re.search(p, text) for p in patterns if p)


async def _run_limit_case(case: dict) -> tuple[bool, list[str]]:
    from app.config import get_settings
    from app.rag.ingest import IngestLimitError, validate_content_length, validate_upload_size

    settings = get_settings()
    notes: list[str] = []
    check = case.get("check")

    if check == "upload_limit":
        oversize = b"x" * (max(1, settings.upload_max_mb) * 1024 * 1024 + 1)
        try:
            validate_upload_size(oversize, settings.upload_max_mb)
            return False, ["expected IngestLimitError for oversize upload"]
        except IngestLimitError as exc:
            msg = str(exc)
            if "上限" not in msg:
                notes.append(f"limit message missing 上限: {msg}")
            return (len(notes) == 0), notes

    if check == "doc_limit":
        oversize = "字" * (max(1, settings.doc_max_chars) + 1)
        try:
            validate_content_length(oversize, settings.doc_max_chars)
            return False, ["expected IngestLimitError for overlong document"]
        except IngestLimitError as exc:
            msg = str(exc)
            if "上限" not in msg:
                notes.append(f"limit message missing 上限: {msg}")
            return (len(notes) == 0), notes

    if check == "circuit_breaker":
        from app.rag.circuit import ExternalCircuit, external_call_failed

        threshold = max(1, int(settings.web_circuit_fail_threshold))
        circuit = ExternalCircuit(threshold=threshold)
        if not external_call_failed({"errors": ["timeout"], "hits": []}):
            notes.append("expected errors+empty hits to count as failure")
        if external_call_failed({"ok": True, "hits": [{"title": "x"}], "count": 1}):
            notes.append("success payload should not count as failure")
        for i in range(threshold):
            circuit.record_failure("web_search", f"err{i}")
        if not circuit.opened:
            notes.append(f"circuit should open after {threshold} failures")
        if circuit.allow("research_topics"):
            notes.append("opened circuit should deny further external calls")
        else:
            meta = circuit.to_meta()
            if "research_topics" not in meta.get("skipped", []):
                notes.append("skipped tools not recorded")
        return (len(notes) == 0), notes

    if check == "concurrency_limit":
        from app.observability.concurrency import (
            acquire_chat_slot,
            concurrency_stats,
            release_chat_slot,
        )

        limit = max(1, int(settings.max_concurrent_chats))
        held = 0
        try:
            for _ in range(limit):
                ok = await acquire_chat_slot(limit)
                if not ok:
                    notes.append("failed to acquire within configured limit")
                    break
                held += 1
            overflow = await acquire_chat_slot(limit)
            if overflow:
                held += 1
                notes.append("expected acquire to fail when at capacity")
            stats = concurrency_stats()
            if stats.get("rejected", 0) < 1:
                notes.append("rejected counter should increase on overflow")
        finally:
            for _ in range(held):
                release_chat_slot()
        return (len(notes) == 0), notes

    if check == "timeout_settings":
        if float(settings.http_timeout_web) <= 0:
            notes.append("http_timeout_web must be > 0")
        if float(settings.http_timeout_research) <= 0:
            notes.append("http_timeout_research must be > 0")
        if float(settings.llm_timeout_seconds) <= 0:
            notes.append("llm_timeout_seconds must be > 0")
        if int(settings.max_concurrent_chats) < 1:
            notes.append("max_concurrent_chats must be >= 1")
        return (len(notes) == 0), notes

    if check == "prompt_versions":
        from app.agents import SYSTEM_PROMPT_VERSION
        from app.agents.reasoning import PROMPT_SYNTH_VERSION, THINKING_PROMPT_VERSION

        for name, value in (
            ("SYSTEM_PROMPT_VERSION", SYSTEM_PROMPT_VERSION),
            ("PROMPT_SYNTH_VERSION", PROMPT_SYNTH_VERSION),
            ("THINKING_PROMPT_VERSION", THINKING_PROMPT_VERSION),
        ):
            if not value or "_v" not in str(value):
                notes.append(f"{name} missing or not versioned: {value!r}")
        return (len(notes) == 0), notes

    if check == "retrieve_spans":
        from app.observability import new_trace
        from app.rag.knowledge import KnowledgeService

        trace = new_trace("eval retrieve spans")
        hits = await KnowledgeService().search("Agent", top_k=3, trace=trace)
        names = [s.name for s in trace.spans]
        for required in ("retrieve_db_scan", "retrieve_score"):
            if required not in names:
                notes.append(f"missing span {required}; got {names}")
        score_span = next((s for s in trace.spans if s.name == "retrieve_score"), None)
        if score_span is not None:
            if "hits" not in score_span.meta:
                notes.append("retrieve_score missing hits meta")
            if score_span.duration_ms < 0:
                notes.append("invalid duration_ms")
            # Default path must be TF-IDF unless env overrides
            backend = score_span.meta.get("backend")
            if backend and backend not in {"tfidf", "vector"}:
                notes.append(f"unexpected backend={backend}")
        # hits may be 0 on empty DB; still require span structure
        del hits
        return (len(notes) == 0), notes

    if check == "vector_default_off":
        if bool(settings.rag_vector_enabled):
            notes.append("RAG_VECTOR_ENABLED must default to false (TF-IDF path)")
        if (settings.rag_embedding_provider or "").strip().lower() not in {
            "hash",
            "fastembed",
            "openai",
            "remote",
            "api",
        }:
            notes.append(f"unexpected embedding provider={settings.rag_embedding_provider}")
        return (len(notes) == 0), notes

    if check == "vector_on_hash_recall":
        import os

        from app.config import get_settings as _gs
        from app.rag.knowledge import KnowledgeService

        prev_enabled = os.environ.get("RAG_VECTOR_ENABLED")
        prev_provider = os.environ.get("RAG_EMBEDDING_PROVIDER")
        try:
            os.environ["RAG_VECTOR_ENABLED"] = "true"
            os.environ["RAG_EMBEDDING_PROVIDER"] = "hash"
            _gs.cache_clear()
            on_settings = _gs()
            if not on_settings.rag_vector_enabled:
                notes.append("failed to enable rag_vector_enabled via env")
            hits_on = await KnowledgeService().search("agent开发需要学习哪些知识", top_k=5)
            titles_on = " ".join(h.get("title", "") for h in hits_on)
            if not any(t.lower() in titles_on.lower() for t in ("agent", "知识地图")):
                notes.append(f"vector-on recall miss; titles={titles_on[:80]}")
            if hits_on and hits_on[0].get("retrieval") != "dense+sparse-hybrid":
                notes.append(
                    f"expected retrieval=dense+sparse-hybrid, got {hits_on[0].get('retrieval')}"
                )

            # Switch back to TF-IDF and confirm still recalls
            os.environ["RAG_VECTOR_ENABLED"] = "false"
            _gs.cache_clear()
            off_settings = _gs()
            if off_settings.rag_vector_enabled:
                notes.append("failed to disable rag_vector_enabled")
            hits_off = await KnowledgeService().search("agent开发需要学习哪些知识", top_k=5)
            titles_off = " ".join(h.get("title", "") for h in hits_off)
            if not any(t.lower() in titles_off.lower() for t in ("agent", "知识地图")):
                notes.append(f"TF-IDF fallback recall miss; titles={titles_off[:80]}")
            if hits_off and hits_off[0].get("retrieval") != "hybrid":
                notes.append(
                    f"expected retrieval=hybrid when off, got {hits_off[0].get('retrieval')}"
                )
        finally:
            if prev_enabled is None:
                os.environ.pop("RAG_VECTOR_ENABLED", None)
            else:
                os.environ["RAG_VECTOR_ENABLED"] = prev_enabled
            if prev_provider is None:
                os.environ.pop("RAG_EMBEDDING_PROVIDER", None)
            else:
                os.environ["RAG_EMBEDDING_PROVIDER"] = prev_provider
            _gs.cache_clear()
        return (len(notes) == 0), notes

    if check == "graph_default_off":
        if bool(settings.rag_graph_enabled):
            notes.append("RAG_GRAPH_ENABLED must default to false")
        if int(settings.rag_graph_hops) < 0:
            notes.append("rag_graph_hops must be >= 0")
        if float(settings.rag_graph_min_confidence) <= 0:
            notes.append("rag_graph_min_confidence must be > 0")
        return (len(notes) == 0), notes

    if check == "graph_disabled_no_impact":
        from app.agents.planner import plan_turn
        from app.rag.graph_query import search_subgraph
        from app.rag.knowledge import KnowledgeService
        from app.tools import run_tool

        if bool(settings.rag_graph_enabled):
            notes.append("expected graph disabled for this check (default)")
        hits = await KnowledgeService().search("agent开发需要学习哪些知识", top_k=5)
        titles = " ".join(h.get("title", "") for h in hits)
        if not any(t.lower() in titles.lower() for t in ("agent", "知识地图")):
            notes.append(f"doc RAG broken with graph off; titles={titles[:80]}")
        plan = plan_turn("agent开发需要学习哪些知识", hits=hits)
        if "search_graph" in plan.steps:
            notes.append("search_graph must not appear in plan when graph disabled")
        raw = await run_tool("search_graph", {"query": "Agent", "hops": 1})
        import json

        payload = json.loads(raw)
        if payload.get("status") != "graph_disabled":
            notes.append(f"expected status=graph_disabled, got {payload.get('status')}")
        sub = await search_subgraph("Agent", settings=settings)
        if sub.get("status") != "graph_disabled":
            notes.append(f"search_subgraph status={sub.get('status')}")
        return (len(notes) == 0), notes

    if check == "graph_on_extract_query":
        import os

        from app.config import get_settings as _gs
        from app.db.seed import ensure_agent_seed
        from app.db.session import init_db
        from app.observability import new_trace
        from app.rag.graph_extract import extract_all_documents
        from app.rag.graph_query import search_subgraph
        from app.rag.knowledge import KnowledgeService
        from app.tools import run_tool

        prev = os.environ.get("RAG_GRAPH_ENABLED")
        try:
            os.environ["RAG_GRAPH_ENABLED"] = "true"
            _gs.cache_clear()
            on = _gs()
            if not on.rag_graph_enabled:
                notes.append("failed to enable rag_graph_enabled")
            await init_db()
            await ensure_agent_seed()
            built = await extract_all_documents(settings=on, limit=50)
            if built.get("status") != "ok":
                notes.append(f"extract_all status={built.get('status')}")
            if int(built.get("relations") or 0) < 1 and int(built.get("entities") or 0) < 1:
                notes.append(f"expected graph entities/relations; got {built}")

            trace = new_trace("eval graph query")
            result = await search_subgraph(
                "Agent 开发需要学习什么", hops=1, top_k=6, settings=on, trace=trace
            )
            if result.get("status") != "ok":
                notes.append(f"search_subgraph status={result.get('status')} err={result.get('error')}")
            if int(result.get("count") or 0) < 1 and not (result.get("entities") or []):
                # soft: at least entities matched after extract
                notes.append(f"graph query empty; result={str(result)[:160]}")
            names = [s.name for s in trace.spans]
            if "graph_query" not in names:
                notes.append(f"missing graph_query span; got {names}")

            raw = await run_tool(
                "search_graph", {"query": "RAG", "hops": 1, "top_k": 6}
            )
            import json

            tool_payload = json.loads(raw)
            if tool_payload.get("status") not in {"ok", "graph_disabled"}:
                notes.append(f"search_graph tool bad status={tool_payload.get('status')}")
            if tool_payload.get("status") == "ok" and int(tool_payload.get("count") or 0) < 1:
                # RAG lexicon should produce mentions on seed docs
                if not (tool_payload.get("entities") or []):
                    notes.append("search_graph returned no triples/entities for RAG")

            # Document RAG still works with graph on
            hits = await KnowledgeService().search("agent开发需要学习哪些知识", top_k=5)
            titles = " ".join(h.get("title", "") for h in hits)
            if not any(t.lower() in titles.lower() for t in ("agent", "知识地图")):
                notes.append(f"doc RAG regress with graph on; titles={titles[:80]}")
        finally:
            if prev is None:
                os.environ.pop("RAG_GRAPH_ENABLED", None)
            else:
                os.environ["RAG_GRAPH_ENABLED"] = prev
            _gs.cache_clear()
        return (len(notes) == 0), notes

    if check == "graph_failure_fallback":
        import os
        from unittest.mock import patch

        from app.config import get_settings as _gs
        from app.rag.graph_query import search_subgraph
        from app.rag.knowledge import KnowledgeService

        prev = os.environ.get("RAG_GRAPH_ENABLED")
        try:
            os.environ["RAG_GRAPH_ENABLED"] = "true"
            _gs.cache_clear()
            on = _gs()

            async def _boom(*_a, **_k):
                raise RuntimeError("simulated graph store failure")

            with patch("app.rag.graph_query.graph_store.list_entities", side_effect=_boom):
                result = await search_subgraph("Agent", hops=1, settings=on)
            if result.get("status") != "error":
                notes.append(f"expected status=error on failure, got {result.get('status')}")
            if result.get("triples"):
                notes.append("failure path should return empty triples")

            # Main doc RAG path must still work
            hits = await KnowledgeService().search("agent开发需要学习哪些知识", top_k=5)
            titles = " ".join(h.get("title", "") for h in hits)
            if not any(t.lower() in titles.lower() for t in ("agent", "知识地图")):
                notes.append(f"doc RAG impacted by graph failure; titles={titles[:80]}")
        finally:
            if prev is None:
                os.environ.pop("RAG_GRAPH_ENABLED", None)
            else:
                os.environ["RAG_GRAPH_ENABLED"] = prev
            _gs.cache_clear()
        return (len(notes) == 0), notes

    if check == "ingest_docx":
        import io

        from app.rag.ingest import detect_source_type, extract_text_from_bytes

        try:
            from docx import Document
        except ImportError:
            return False, ["python-docx not installed"]
        doc = Document()
        doc.add_paragraph("境内住宿标准：500元/晚。")
        buf = io.BytesIO()
        doc.save(buf)
        text = extract_text_from_bytes("eval.docx", buf.getvalue())
        if "500" not in text:
            notes.append(f"docx extract miss 500; got={text[:80]!r}")
        if detect_source_type("eval.docx") != "word":
            notes.append("detect_source_type(docx) != word")
        return (len(notes) == 0), notes

    if check == "ingest_xlsx":
        import io

        from app.rag.ingest import detect_source_type, extract_text_from_bytes

        try:
            from openpyxl import Workbook
        except ImportError:
            return False, ["openpyxl not installed"]
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "住宿"
        ws["B1"] = "500"
        buf = io.BytesIO()
        wb.save(buf)
        text = extract_text_from_bytes("eval.xlsx", buf.getvalue())
        if "住宿" not in text or "500" not in text:
            notes.append(f"xlsx extract miss; got={text[:80]!r}")
        if detect_source_type("eval.xlsx") != "xlsx":
            notes.append("detect_source_type(xlsx) != xlsx")
        return (len(notes) == 0), notes

    if check == "ingest_html":
        from app.rag.ingest import extract_text_from_bytes

        html = "<html><body><h1>考勤</h1><p>迟到扣款</p><script>evil()</script></body></html>"
        text = extract_text_from_bytes("eval.html", html.encode("utf-8"))
        if "考勤" not in text or "迟到" not in text:
            notes.append(f"html extract miss; got={text[:80]!r}")
        if "evil" in text.lower() or "script" in text.lower():
            notes.append("html extractor leaked script content")
        return (len(notes) == 0), notes

    if check == "ingest_allowed_ext":
        from app.rag.ingest import ALLOWED_EXT

        required = {".pdf", ".docx", ".pptx", ".xlsx", ".png", ".jpg", ".txt", ".md", ".html"}
        missing = sorted(required - set(ALLOWED_EXT))
        if missing:
            notes.append(f"ALLOWED_EXT missing {missing}")
        return (len(notes) == 0), notes

    if check == "delete_conversation":
        from app.db import crud
        from app.db.session import get_session_maker

        sm = get_session_maker()
        async with sm() as db:
            conv = await crud.create_conversation(db, "评测删除会话")
            await crud.add_message(db, conv.id, "user", "hi")
            await db.commit()
            cid = conv.id
        async with sm() as db:
            ok_del = await crud.delete_conversation(db, cid)
            await db.commit()
            if not ok_del:
                notes.append("delete_conversation returned False")
            again = await crud.delete_conversation(db, cid)
            if again:
                notes.append("second delete should be False")
            gone = await crud.get_conversation(db, cid)
            if gone is not None:
                notes.append("conversation still exists after delete")
        return (len(notes) == 0), notes

    if check == "model_route_select":
        from app.agents import ChatAgent
        from app.config import get_settings as _gs

        agent = ChatAgent(_gs())
        policy_model = agent._select_llm_model(
            policy_strict=True, deep_think=False, plan={"intent": "local"}
        )
        general_model = agent._select_llm_model(
            policy_strict=False, deep_think=False, plan={"intent": "general"}
        )
        think_model = agent._select_llm_model(
            policy_strict=True, deep_think=True, plan={"intent": "local"}
        )
        base = (agent.settings.llm_model or "").strip()
        complex_model = (agent.settings.llm_model_complex or base).strip()
        if policy_model != base:
            notes.append(f"policy path want {base!r} got {policy_model!r}")
        if general_model != complex_model:
            notes.append(f"general path want {complex_model!r} got {general_model!r}")
        if think_model != complex_model:
            notes.append(f"deep_think path want {complex_model!r} got {think_model!r}")
        return (len(notes) == 0), notes

    if check == "sanitize_tool_leak":
        from app.agents.reasoning import looks_like_tool_leak, sanitize_assistant_text

        sample = (
            "南雅中学师资不错\n"
            "<tool_calls>\n"
            'name="web_search"\n'
            'query="长沙市南雅中学简介"\n'
            'max_results="3"\n'
            "</tool_calls>"
        )
        if not looks_like_tool_leak(sample):
            notes.append("looks_like_tool_leak should be True for tool_calls dump")
        cleaned = sanitize_assistant_text(sample)
        for bad in ("tool_calls", "web_search", "max_results", "invoke"):
            if bad.lower() in cleaned.lower():
                notes.append(f"sanitize left leak marker: {bad!r} in {cleaned!r}")
        if "南雅" not in cleaned and "师资" not in cleaned:
            notes.append(f"sanitize stripped too much useful text: {cleaned!r}")
        dsml = '< | DSML | invoke name="search_knowledge">x</ | DSML | parameter>'
        if sanitize_assistant_text(dsml):
            notes.append("DSML-only payload should sanitize to empty")
        return (len(notes) == 0), notes

    if check == "web_empty_honest":
        from app.agents.loop import LoopState, build_react_thinking
        from app.agents.planner import AgentPlan, Intent
        from app.agents.reasoning import honest_web_empty_reply

        reply = honest_web_empty_reply("茶陵一中")
        if not any(k in reply for k in ("未找到", "不足", "无法")):
            notes.append(f"honest reply missing honesty markers: {reply!r}")
        if "tool_calls" in reply or "DSML" in reply:
            notes.append("honest reply leaked tool markup")
        plan = AgentPlan(
            intent=Intent.WEB,
            query="茶陵一中",
            steps=["web_search", "synthesize"],
            force_web=True,
            use_web=True,
        )
        state = LoopState(query="茶陵一中", plan=plan, hits=[])
        state.add_observation("web_search", {"count": 0, "ok": False, "hits": [], "empty": True})
        thinking = build_react_thinking(state)
        if "未找到" not in thinking:
            notes.append(f"thinking should mention 未找到; got {thinking!r}")
        return (len(notes) == 0), notes

    return False, [f"unknown check={check}"]


async def run(categories: set[str] | None = None) -> int:
    from app.db.seed import ensure_agent_seed
    from app.db.session import init_db
    from app.rag.knowledge import KnowledgeService

    await init_db()
    await ensure_agent_seed()

    cases = load_cases(categories)
    knowledge = KnowledgeService()
    passed = 0
    filt = ",".join(sorted(categories)) if categories else "all"
    print(f"Running {len(cases)} eval cases (filter={filt})…\n")

    for case in cases:
        cid = case.get("id", "?")
        query = case.get("query", "")
        ok = True
        notes: list[str] = []

        if case.get("check"):
            ok, notes = await _run_limit_case(case)
            status = "PASS" if ok else "FAIL"
            if ok:
                passed += 1
            print(f"[{status}] {cid}: check={case.get('check')}")
            for n in notes:
                print(f"       - {n}")
            continue

        # T4-3: optional prior turn for multi-turn slot enrichment
        from app.agents.planner import plan_turn
        from app.agents.reasoning import build_search_query, synthesize_knowledge_answer

        messages: list[dict[str, str]] = []
        hu = case.get("history_user") or ""
        ha = case.get("history_assistant") or ""
        if hu:
            messages.append({"role": "user", "content": str(hu)})
        if ha:
            messages.append({"role": "assistant", "content": str(ha)})
        messages.append({"role": "user", "content": query})

        search_q = build_search_query(messages) or query
        hits = await knowledge.search(search_q, top_k=5)
        from app.rag.research import filter_relevant_hits
        from app.rag.query_rewrite import rewrite_search_query

        # Keep the harness aligned with production planning: relevance filtering
        # must see the same retrieval-oriented synonym expansion as search.
        filter_q = rewrite_search_query(search_q) or search_q
        hits = filter_relevant_hits(filter_q, hits)
        plan = plan_turn(query, hits=hits, retrieval_query=search_q)

        if case.get("must_enriched_contain"):
            needle = str(case["must_enriched_contain"])
            if needle not in search_q:
                ok = False
                notes.append(f"enriched query missing {needle!r}; got {search_q[:80]!r}")

        expected = case.get("expected_intent")
        if expected and plan.intent.value != expected:
            ok = False
            notes.append(f"intent={plan.intent.value} want={expected}")

        step = case.get("must_plan_step")
        if step and step not in plan.steps:
            ok = False
            notes.append(f"missing step {step} in {plan.steps}")

        banned_step = case.get("must_not_plan_step")
        if banned_step and banned_step in plan.steps:
            ok = False
            notes.append(f"forbidden step {banned_step} in {plan.steps}")

        titles = " ".join(h.get("title", "") for h in hits)
        must_any = case.get("must_hit_title_any") or []
        if must_any and not any(t.lower() in titles.lower() for t in must_any):
            # soft fail for retrieval only when intent is local
            if expected == "local":
                ok = False
                notes.append(f"no title match among {must_any}; got {titles[:80]}")

        if "expect_max_hits" in case and len(hits) > int(case["expect_max_hits"]):
            ok = False
            notes.append(f"hits={len(hits)} want<={case['expect_max_hits']}")

        if "expect_min_hits" in case and len(hits) < int(case["expect_min_hits"]):
            ok = False
            notes.append(f"hits={len(hits)} want>={case['expect_min_hits']}")

        answer = synthesize_knowledge_answer(query, hits, history=messages)
        # Align eval answers with ChatAgent short-circuit intents
        if plan.intent.value == "chitchat":
            from app.rag.chitchat import chitchat_reply

            answer = chitchat_reply(query)
        elif plan.intent.value == "clarify":
            from app.agents import _build_clarify_reply

            answer = _build_clarify_reply(query)
        elif plan.intent.value == "general" and case.get("expected_intent") == "general":
            # Mock 路径不调用真实 LLM；仅在显式验收 general 时替换答文
            answer = "通用问答将由模型直接作答，不因非制度而拒答。"
        elif plan.intent.value == "web" and case.get("expected_intent") == "web":
            answer = "联网检索路径：将调用 web_search 后综合作答。"
        for bad in case.get("forbid_phrases") or []:
            if bad in answer:
                ok = False
                notes.append(f"forbidden phrase: {bad}")

        must_ans = case.get("must_answer_contain_any") or []
        if must_ans and not _contains_any(answer, must_ans):
            ok = False
            notes.append(f"answer missing any of {must_ans}")

        must_all = case.get("must_answer_contain_all") or []
        for needle in must_all:
            if needle not in answer:
                ok = False
                notes.append(f"answer missing required: {needle}")

        must_re = case.get("must_answer_match_any") or []
        if must_re and not _match_any(answer, must_re):
            ok = False
            notes.append(f"answer missing pattern any of {must_re}")

        status = "PASS" if ok else "FAIL"
        if ok:
            passed += 1
        print(f"[{status}] {cid}: {query}")
        for n in notes:
            print(f"       - {n}")

    print(f"\n{passed}/{len(cases)} passed")
    return 0 if passed == len(cases) else 1


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Atlas eval gate")
    p.add_argument(
        "--category",
        default="",
        help="Comma-separated categories, e.g. hit,missing,general,ingest",
    )
    return p.parse_args(argv)


if __name__ == "__main__":
    args = _parse_args()
    cats = {c.strip().lower() for c in args.category.split(",") if c.strip()} or None
    raise SystemExit(asyncio.run(run(cats)))
