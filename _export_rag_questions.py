# -*- coding: utf-8 -*-
from pathlib import Path
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

DOCX = Path(r"c:\Users\czy2004\Desktop\RAG面试追问问题清单.docx")
PDF = Path(r"c:\Users\czy2004\Desktop\RAG面试追问问题清单.pdf")
PS1 = Path(r"c:\Users\czy2004\Desktop\AI 全栈开发\_export_qa_pdf.ps1")


def font(run, name="宋体", size=11, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), name)


def p(doc, text, *, size=11, bold=False, name="宋体", align=None, after=6, first=False):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    pf = para.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first:
        pf.first_line_indent = Pt(22)
    run = para.add_run(text)
    font(run, name=name, size=size, bold=bold)
    return para


def h(doc, text, level=1):
    if level == 1:
        return p(doc, text, size=16, bold=True, name="黑体", after=10)
    if level == 2:
        return p(doc, text, size=13, bold=True, name="黑体", after=8)
    return p(doc, text, size=11, bold=True, name="黑体", after=6)


def bullets(doc, items):
    for t in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(t)
        font(run, size=11)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    p(doc, "Atlas / RAG 知识助手", size=12, bold=True, name="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    p(doc, "长图截图问题提取清单", size=18, bold=True, name="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    p(
        doc,
        "来源：微信对话长图（多轮检索 / 模糊问法 / 拆问并行 / 同义召回 / 空库 / 超时）",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=14,
    )
    p(doc, "说明：以下问题从聊天截图中按原意整理，按主题归类，便于答辩准备与方案对照。", first=True)

    h(doc, "一、多轮上下文与检索策略（截图 1）", 1)
    p(
        doc,
        "场景铺垫（对方原话大意）：用户第一句「查询员工制度」，第二句「那薪资呢？」——"
        "不能只把「那薪资呢」四个字拿去切块检索。",
        first=True,
    )
    h(doc, "1.1 核心追问", 2)
    bullets(
        doc,
        [
            "第一句问「查询员工制度」，系统能不能正常检索出来？",
            "第二句「那薪资呢？」能不能拼接上下文？有没有记忆？",
            "会不会直接把「那薪资呢」这四个字切块去查？（对方认为这显然不合理）",
            "你们是怎么查的？查询策略具体是什么？",
            "是不是直接「切块 → 算相似度 → 召回 → rerank」就够了？",
            "如果这一次检索效果不好，怎么办？",
        ],
    )
    h(doc, "1.2 可写成答辩要点的标准问法", 2)
    bullets(
        doc,
        [
            "多轮追问时，检索 query 如何改写 / 槽位补全？",
            "标准 RAG 流水线是否足够？还缺哪些环节（改写、过滤、拒答等）？",
            "检索质量差时的降级与兜底策略是什么？",
        ],
    )

    h(doc, "二、模糊问、复杂问、同义与异常（截图 2）", 1)
    h(doc, "2.1 模糊 / 极短问", 2)
    bullets(
        doc,
        [
            "如果用户问得很模糊怎么办？",
            "例如只问两个字：「薪资」——怎么处理？",
            "有没有反馈机制，让用户补充关键信息？",
        ],
    )
    h(doc, "2.2 复合问题与并行检索", 2)
    bullets(
        doc,
        [
            "如果用户问：「薪资和员工制度如何，社招和校招有什么区别」——怎么处理？",
            "能不能把这句话拆成几个子问题，并行分别查询，再汇总回答？",
        ],
    )
    h(doc, "2.3 同义词 / 别名召回", 2)
    bullets(
        doc,
        [
            "知识库里是 aaa，用户问 bbb，但 aaa 与 bbb 其实是同一事物的两种说法——能不能查到？",
        ],
    )
    h(doc, "2.4 空库与系统可靠性", 2)
    bullets(
        doc,
        [
            "如果知识库里没有相关内容，应该怎么回答？",
            "如果 Agent 的 API 卡住、请求超时，怎么办？",
        ],
    )

    h(doc, "三、问题总表（便于勾选准备）", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, t in enumerate(["编号", "主题", "问题摘要", "考查点"]):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(t)
        font(run, name="黑体", size=9, bold=True)

    rows = [
        ("Q1", "多轮记忆", "「那薪资呢」能否带上「员工制度」上下文", "追问改写 / 槽位 / 记忆"),
        ("Q2", "检索策略", "怎么查？策略是什么", "rewrite + hybrid + 过滤"),
        ("Q3", "流水线", "切块+相似度+召回+rerank 是否够", "生产级 RAG 完备性"),
        ("Q4", "差检索", "检索效果不好怎么办", "阈值门禁 / 拒答 / 重试"),
        ("Q5", "短问澄清", "只问「薪资」是否反问补全", "澄清对话 / 缺槽提示"),
        ("Q6", "问题分解", "复合问能否拆子问题并行查再汇总", "query decompose"),
        ("Q7", "同义映射", "aaa/bbb 别名能否召回", "同义词 / 向量语义"),
        ("Q8", "空库", "库里没有怎么答", "拒答不编造"),
        ("Q9", "超时", "API 卡住/超时怎么办", "timeout / 降级 / 熔断"),
    ]
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            font(run, size=9)

    doc.add_paragraph()
    p(doc, "表 1　截图问题总览", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    h(doc, "四、附：界面截图中出现的示例用户问句（非面试追问）", 1)
    bullets(
        doc,
        [
            "联网：今天有哪些科技新闻",
            "差旅住宿费上限是多少？（输入框示例）",
            "agent开发需要学习哪些知识",
            "FastAPI 是什么？",
            "北师香港浸会大学就业去向",
            "展开第2点 / 为什么需要 RAG（系统提示的追问示例）",
        ],
    )
    p(
        doc,
        "整理说明：面试追问以「一、二、三」为主；第四节仅为产品界面里出现过的问句，便于区分。",
        size=10,
        first=True,
        after=8,
    )
    p(doc, "文档由聊天长图识别整理 · Atlas RAG 答辩备问", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.save(str(DOCX))


def export_pdf():
    PS1.write_text(
        f"""$docx = '{DOCX}'
$pdf = '{PDF}'
$app = New-Object -ComObject Word.Application
$app.Visible = $false
$d = $app.Documents.Open($docx)
$d.SaveAs([ref]$pdf, [ref]17)
$d.Close()
$app.Quit()
""",
        encoding="utf-8",
    )
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(PS1)],
        check=True,
    )


if __name__ == "__main__":
    build()
    export_pdf()
    print("DOCX", DOCX, DOCX.stat().st_size)
    print("PDF", PDF, PDF.stat().st_size)
